import streamlit as st
import pandas as pd
import os
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import column_index_from_string
from openpyxl.cell.cell import MergedCell

# ----------------- CONFIG APP -----------------

st.set_page_config(
    page_title="FT Grand Compte",
    page_icon="🚚",
    layout="wide"
)

st.title("Generateur de Fiches Techniques Grand Compte")
st.caption("Version de test basée sur bdd_CG.xlsx")

IMG_ROOT = "images"  # dossier racine des images (dans le repo)


# ----------------- FONCTIONS UTILES -----------------

def resolve_image_path(cell_value, subdir):
    """
    Transforme la valeur Excel (chemin complet, nom de fichier ou URL)
    en chemin exploitable.
    - subdir : sous-dossier dans IMG_ROOT (ex : 'vehicules', 'clients', 'carburant')
    """
    if not isinstance(cell_value, str) or not cell_value.strip():
        return None

    val = cell_value.strip()

    # URL
    if val.lower().startswith(("http://", "https://")):
        return val

    # Chemin local / nom de fichier
    filename = os.path.basename(val)
    return os.path.join(IMG_ROOT, subdir, filename)


def show_image(path_or_url, caption, uploaded_file=None):
    """
    Affiche une image :
      - si uploaded_file est fourni -> on affiche l'upload
      - sinon on affiche path_or_url (URL ou chemin local)
    """
    st.caption(caption)

    if uploaded_file is not None:
        st.image(uploaded_file)
        return

    if not path_or_url:
        st.info("Pas d'image définie")
        return

    # URL
    if isinstance(path_or_url, str) and path_or_url.lower().startswith(("http://", "https://")):
        st.image(path_or_url)
        return

    # Fichier local
    if os.path.exists(path_or_url):
        st.image(path_or_url)
    else:
        st.warning(f"Image introuvable : {path_or_url}")


@st.cache_data
def load_data():
    xls = pd.ExcelFile("bdd_CG.xlsx")

    vehicules = pd.read_excel(xls, "FS_referentiel_produits_std_Ver")
    cabines = pd.read_excel(xls, "CABINES")
    moteurs = pd.read_excel(xls, "MOTEURS")
    chassis = pd.read_excel(xls, "CHASSIS")
    caisses = pd.read_excel(xls, "CAISSES")
    frigo = pd.read_excel(xls, "FRIGO")
    hayons = pd.read_excel(xls, "HAYONS")

    return vehicules, cabines, moteurs, chassis, caisses, frigo, hayons


def filtre_select(df, col, label):
    """Selectbox simple dans la sidebar sur une colonne (avec 'Tous')."""
    if col not in df.columns:
        st.sidebar.write(f"(colonne '{col}' absente)")
        return df, None

    vals = df[col].dropna()
    vals = [v for v in vals.unique().tolist() if str(v).strip() != ""]
    vals = sorted(vals)

    options_display = ["Tous"] + vals
    choix = st.sidebar.selectbox(label, options_display)

    if choix != "Tous":
        df = df[df[col] == choix]

    return df, choix


def format_vehicule(row):
    champs = []
    for c in ["code_pays", "Marque", "Modele", "Code_PF", "Standard_PF"]:
        if c in row and pd.notna(row[c]):
            champs.append(str(row[c]))
    return " – ".join(champs)


# ----------------- GENERATION DE LA FT -----------------

def genere_ft_excel(
    veh,
    cab_prod_choice, cab_opt_choice,
    mot_prod_choice, mot_opt_choice,
    ch_prod_choice, ch_opt_choice,
    caisse_prod_choice, caisse_opt_choice,
    gf_prod_choice, gf_opt_choice,
    hay_prod_choice, hay_opt_choice,
    img_veh_upload=None,
    img_client_upload=None,
    img_carbu_upload=None,
):
    """
    Génère une fiche technique Excel à partir du modèle 'FT_Grand_Compte.xlsx'
    en utilisant :
      - la ligne véhicule sélectionnée (veh)
      - les choix de codes produit / options pour chaque composant
      - éventuellement des images uploadées (véhicule / client / carburant)
    """

    global cabines, moteurs, chassis, caisses, frigo, hayons

    template_path = "FT_Grand_Compte.xlsx"
    if not os.path.exists(template_path):
        st.error("Modèle FT_Grand_Compte.xlsx manquant dans le dossier.")
        return None

    wb = load_workbook(template_path, read_only=False, data_only=False)
    ws = wb["date"]  # nom de l’onglet dans le modèle

    # ---------- petits helpers internes ----------

    def build_values(row, code_col):
        """Transforme une ligne de BDD composant en liste de valeurs (sans libellé)."""
        if row is None:
            return []
        vals = []
        for col, val in row.items():
            if col == code_col:
                continue
            if pd.isna(val) or str(val).strip() == "":
                continue
            name_lower = str(col).lower()
            if "produit (p" in name_lower and "option" in name_lower:
                continue
            if name_lower.startswith("zone libre"):
                continue
            if col == "_":
                continue
            vals.append(str(val))
        return vals

    def write_block(start_cell, values, max_rows):
        """Écrit les valeurs ligne par ligne à partir de start_cell (1 seule colonne)."""
        if max_rows <= 0:
            return
        col_letters = "".join(ch for ch in start_cell if ch.isalpha())
        row_digits = "".join(ch for ch in start_cell if ch.isdigit())
        start_col = column_index_from_string(col_letters)
        start_row = int(row_digits)

        for i in range(max_rows):
            cell = ws.cell(row=start_row + i, column=start_col)
            if isinstance(cell, MergedCell):
                continue
            if i < len(values):
                cell.value = values[i]
            else:
                cell.value = None

    def find_row(df, code, code_col):
        """Retourne la première ligne de df où df[code_col] == code."""
        if not isinstance(code, str) or code.strip() == "":
            return None
        cand = df[df[code_col] == code]
        if cand.empty:
            return None
        return cand.iloc[0]

    def choose_codes(prod_choice, opt_choice, veh_prod, veh_opt):
        """
        A partir des 2 filtres (code produit et code option),
        renvoie (code_produit, code_option) à utiliser.
        """
        prod_code = veh_prod
        opt_code = veh_opt
        if isinstance(prod_choice, str) and prod_choice not in (None, "", "Tous"):
            prod_code = prod_choice
        if isinstance(opt_choice, str) and opt_choice not in (None, "", "Tous"):
            opt_code = opt_choice
        return prod_code, opt_code

    def set_row_text(row_idx, search_substring, new_text):
        """
        Cherche dans la ligne row_idx la cellule qui contient search_substring
        et remplace son contenu par new_text (utile pour les cellules fusionnées).
        """
        for cell in ws[row_idx]:
            if isinstance(cell, MergedCell):
                continue
            val = cell.value
            if isinstance(val, str) and search_substring.lower() in val.lower():
                cell.value = new_text
                break

    def place_image(img_obj, anchor, max_w=None, max_h=None):
        """Redimensionne et place une image à une position donnée."""
        if img_obj is None:
            return
        w, h = img_obj.width, img_obj.height
        ratio = 1.0
        if max_w and w > max_w:
            ratio = min(ratio, max_w / w)
        if max_h and h > max_h:
            ratio = min(ratio, max_h / h)
        img_obj.width = int(w * ratio)
        img_obj.height = int(h * ratio)
        img_obj.anchor = anchor
        ws.add_image(img_obj)

    def set_merged_value(ws_local, cell_ref, value):
        """
        Écrit `value` dans cell_ref.
        Si cell_ref est une MergedCell, on écrit dans la cellule
        en haut à gauche de la plage fusionnée correspondante.
        """
        cell = ws_local[cell_ref]

        # Si ce n'est pas une cellule fusionnée -> écriture directe
        if not isinstance(cell, MergedCell):
            cell.value = value
            return

        # Sinon on retrouve la plage fusionnée contenant la cellule
        col_letters = "".join(ch for ch in cell_ref if ch.isalpha())
        row_digits = "".join(ch for ch in cell_ref if ch.isdigit())
        col_idx = column_index_from_string(col_letters)
        row_idx = int(row_digits)

        for merged_range in ws_local.merged_cells.ranges:
            if (
                merged_range.min_row <= row_idx <= merged_range.max_row
                and merged_range.min_col <= col_idx <= merged_range.max_col
            ):
                top_left = ws_local.cell(row=merged_range.min_row, column=merged_range.min_col)
                top_left.value = value
                return

        # Si aucune plage fusionnée ne correspond (cas rare)
        cell.value = value

    # ---------- 1) En-tête "classique" (colonnes C5..C12) ----------

    mapping_header = {
        "code_pays": "C5",
        "Marque": "C6",
        "Modele": "C7",
        "Code_PF": "C8",
        "Standard_PF": "C9",
        "catalogue_1\n PF": "C10",
        "catalogue_2\nST": "C11",
        "catalogue_3\n LIBRE": "C12",
    }
    for col_bdd, cell_addr in mapping_header.items():
        if col_bdd in veh.index:
            ws[cell_addr] = veh[col_bdd]

    # ---------- 1bis) Lignes 1 et 2 : titre + Code PF ----------

    # Combinaison véhicule (comme dans le select)
    veh_parts = []
    for key in ["code_pays", "Marque", "Modele", "Code_PF", "Standard_PF"]:
        if key in veh.index and pd.notna(veh[key]):
            veh_parts.append(str(veh[key]))
    comb_veh = " - ".join(veh_parts)

    # Type de caisse : frigo si groupe frigo renseigné, sinon sèche
    gf_val = veh.get("C_Groupe frigo")
    if isinstance(gf_val, str) and gf_val.strip() and gf_val.strip().upper() != "GF_VIDE":
        type_caisse = "CAISSE FRIGORIFIQUE"
    else:
        type_caisse = "CAISSE SECHE"

    titre_ligne1 = f"{comb_veh} - {type_caisse}"
    code_pf_val = str(veh.get("Code_PF", "") or "")

    # On écrit dans le bandeau vert (zone fusionnée centrée sur la ligne)
    # Ici on vise C1 et C2, qui appartiennent au bandeau vert
    set_merged_value(ws, "C1", titre_ligne1)
    set_merged_value(ws, "C2", f"CODE PF : {code_pf_val}")

    # ---------- 2) Images (véhicule, client, carburant) ----------

    # Positions dans le modèle (à ajuster si besoin)
    VEH_ANCHOR = "B7"     # véhicule au centre
    CLIENT_ANCHOR = "F8"  # logo client
    CARBU_ANCHOR = "F11"  # picto carburant

    # Tailles max
    VEH_MAX_W, VEH_MAX_H = 800, 520
    LOGO_MAX_W, LOGO_MAX_H = 320, 240
    CARBU_MAX_W, CARBU_MAX_H = 260, 220

    # Chemins issus de la BDD si pas d’upload
    img_veh_path = resolve_image_path(veh.get("Image Vehicule"), "vehicules")
    img_client_path = resolve_image_path(veh.get("Image Client"), "clients")
    img_carbu_path = resolve_image_path(veh.get("Image Carburant"), "carburant")

    # Image véhicule
    if img_veh_upload is not None:
        data = img_veh_upload.read()
        img_veh_upload.seek(0)
        place_image(XLImage(BytesIO(data)), VEH_ANCHOR, VEH_MAX_W, VEH_MAX_H)
    elif img_veh_path and os.path.exists(img_veh_path):
        place_image(XLImage(img_veh_path), VEH_ANCHOR, VEH_MAX_W, VEH_MAX_H)

    # Logo client
    if img_client_upload is not None:
        data = img_client_upload.read()
        img_client_upload.seek(0)
        place_image(XLImage(BytesIO(data)), CLIENT_ANCHOR, LOGO_MAX_W, LOGO_MAX_H)
    elif img_client_path and os.path.exists(img_client_path):
        place_image(XLImage(img_client_path), CLIENT_ANCHOR, LOGO_MAX_W, LOGO_MAX_H)

    # Picto carburant
    if img_carbu_upload is not None:
        data = img_carbu_upload.read()
        img_carbu_upload.seek(0)
        place_image(XLImage(BytesIO(data)), CARBU_ANCHOR, CARBU_MAX_W, CARBU_MAX_H)
    elif img_carbu_path and os.path.exists(img_carbu_path):
        place_image(XLImage(img_carbu_path), CARBU_ANCHOR, CARBU_MAX_W, CARBU_MAX_H)

    # ---------- 3) Composants & options ----------

    # CABINE
    cab_prod_code, cab_opt_code = choose_codes(
        cab_prod_choice, cab_opt_choice,
        veh.get("C_Cabine"),
        veh.get("C_Cabine-OPTIONS"),
    )
    cab_prod_row = find_row(cabines, cab_prod_code, "C_Cabine")
    cab_opt_row = find_row(cabines, cab_opt_code, "C_Cabine")
    write_block("B18", build_values(cab_prod_row, "C_Cabine"), max_rows=17)
    write_block("B37", build_values(cab_opt_row, "C_Cabine"), max_rows=3)

    # MOTEUR
    mot_prod_code, mot_opt_code = choose_codes(
        mot_prod_choice, mot_opt_choice,
        veh.get("M_moteur"),
        veh.get("M_moteur-OPTIONS"),
    )
    mot_prod_row = find_row(moteurs, mot_prod_code, "M_moteur")
    mot_opt_row = find_row(moteurs, mot_opt_code, "M_moteur")
    write_block("F18", build_values(mot_prod_row, "M_moteur"), max_rows=17)
    write_block("F37", build_values(mot_opt_row, "M_moteur"), max_rows=3)

    # CHASSIS
    ch_prod_code, ch_opt_code = choose_codes(
        ch_prod_choice, ch_opt_choice,
        veh.get("C_Chassis"),
        veh.get("C_Chassis-OPTIONS"),
    )
    ch_prod_row = find_row(chassis, ch_prod_code, "c_chassis")
    ch_opt_row = find_row(chassis, ch_opt_code, "c_chassis")
    write_block("H18", build_values(ch_prod_row, "c_chassis"), max_rows=17)
    write_block("H37", build_values(ch_opt_row, "c_chassis"), max_rows=3)

    # CAISSE / CARROSSERIE
    caisse_prod_code, caisse_opt_code = choose_codes(
        caisse_prod_choice, caisse_opt_choice,
        veh.get("C_Caisse"),
        veh.get("C_Caisse-OPTIONS"),
    )
    caisse_prod_row = find_row(caisses, caisse_prod_code, "c_caisse")
    caisse_opt_row = find_row(caisses, caisse_opt_code, "c_caisse")
    write_block("B40", build_values(caisse_prod_row, "c_caisse"), max_rows=5)
    write_block("B47", build_values(caisse_opt_row, "c_caisse"), max_rows=2)

    # GROUPE FRIGO
    gf_prod_code, gf_opt_code = choose_codes(
        gf_prod_choice, gf_opt_choice,
        veh.get("C_Groupe frigo"),
        veh.get("C_Groupe frigo-OPTIONS"),
    )
    gf_prod_row = find_row(frigo, gf_prod_code, "c_groupe frigo")
    gf_opt_row = find_row(frigo, gf_opt_code, "c_groupe frigo")
    write_block("B51", build_values(gf_prod_row, "c_groupe frigo"), max_rows=6)
    write_block("B59", build_values(gf_opt_row, "c_groupe frigo"), max_rows=2)

    # HAYON
    hay_prod_code, hay_opt_code = choose_codes(
        hay_prod_choice, hay_opt_choice,
        veh.get("C_Hayon elevateur"),
        veh.get("C_Hayon elevateur-OPTIONS"),
    )
    hay_prod_row = find_row(hayons, hay_prod_code, "c_hayon elevateur")
    hay_opt_row = find_row(hayons, hay_opt_code, "c_hayon elevateur")
    write_block("B61", build_values(hay_prod_row, "c_hayon elevateur"), max_rows=5)
    write_block("B68", build_values(hay_opt_row, "c_hayon elevateur"), max_rows=3)

    # ---------- 4) Dimensions & poids ----------

    col_Wint   = "W int\n utile \nsur plinthe"
    col_Lint   = "L int \nutile \nsur plinthe"
    col_Hint   = "H int"
    col_Hhors  = "H"
    col_L      = "L"
    col_Z      = "Z"
    col_Hc     = "Hc"
    col_F      = "F"
    col_X      = "X"
    col_pal    = "palettes 800 x 1200 mm"
    col_PTAC   = "PTAC"
    col_CU     = "CU"
    col_volume = "Volume"

    ws["I5"]  = veh.get(col_Wint)
    ws["I6"]  = veh.get(col_Lint)
    ws["I7"]  = veh.get(col_Hint)
    ws["I8"]  = veh.get(col_Hhors)

    ws["K4"]  = veh.get(col_L)
    ws["K5"]  = veh.get(col_Z)
    ws["K6"]  = veh.get(col_Hc)
    ws["K7"]  = veh.get(col_F)
    ws["K8"]  = veh.get(col_X)

    ws["I10"] = veh.get(col_PTAC)
    ws["I11"] = veh.get(col_CU)
    ws["I12"] = veh.get(col_volume)
    ws["I13"] = veh.get(col_pal)

    # ---------- 5) Sauvegarde dans un buffer ----------

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ----------------- CHARGEMENT DES DONNÉES -----------------

vehicules, cabines, moteurs, chassis, caisses, frigo, hayons = load_data()


# ----------------- SIDEBAR : FILTRES -----------------

st.sidebar.header("Filtres véhicule")

df_filtre = vehicules.copy()

# Filtres principaux
df_filtre, code_pays = filtre_select(df_filtre, "code_pays", "Code pays")
df_filtre, marque = filtre_select(df_filtre, "Marque", "Marque")
df_filtre, modele = filtre_select(df_filtre, "Modele", "Modèle")
df_filtre, code_pf = filtre_select(df_filtre, "Code_PF", "Code PF")
df_filtre, std_pf = filtre_select(df_filtre, "Standard_PF", "Standard PF")

# Filtres composants : 1 filtre code produit + 1 filtre code options

df_filtre, cab_prod_choice = filtre_select(df_filtre, "C_Cabine", "Cabine - code produit")
df_filtre, cab_opt_choice = filtre_select(df_filtre, "C_Cabine-OPTIONS", "Cabine - code options")

df_filtre, mot_prod_choice = filtre_select(df_filtre, "M_moteur", "Moteur - code produit")
df_filtre, mot_opt_choice = filtre_select(df_filtre, "M_moteur-OPTIONS", "Moteur - code options")

df_filtre, ch_prod_choice = filtre_select(df_filtre, "C_Chassis", "Châssis - code produit")
df_filtre, ch_opt_choice = filtre_select(df_filtre, "C_Chassis-OPTIONS", "Châssis - code options")

df_filtre, caisse_prod_choice = filtre_select(df_filtre, "C_Caisse", "Caisse - code produit")
df_filtre, caisse_opt_choice = filtre_select(df_filtre, "C_Caisse-OPTIONS", "Caisse - code options")

df_filtre, gf_prod_choice = filtre_select(df_filtre, "C_Groupe frigo", "Groupe frigo - code produit")
df_filtre, gf_opt_choice = filtre_select(df_filtre, "C_Groupe frigo-OPTIONS", "Groupe frigo - code options")

df_filtre, hay_prod_choice = filtre_select(df_filtre, "C_Hayon elevateur", "Hayon - code produit")
df_filtre, hay_opt_choice = filtre_select(df_filtre, "C_Hayon elevateur-OPTIONS", "Hayon - code options")

st.subheader("Résultats du filtrage")
st.write(f"{len(df_filtre)} combinaison(s) véhicule trouvée(s).")

if df_filtre.empty:
    st.warning("Aucun véhicule ne correspond aux filtres sélectionnés.")
    st.stop()

# ----------------- CHOIX D'UNE LIGNE -----------------

indices = list(df_filtre.index)
choix_idx = st.selectbox(
    "Sélectionne le véhicule pour générer la FT :",
    indices,
    format_func=lambda i: format_vehicule(df_filtre.loc[i]),
)

veh = df_filtre.loc[choix_idx]

st.markdown("---")
st.subheader("Synthèse véhicule")

cols_synthese = [
    "code_pays",
    "Marque",
    "Modele",
    "Code_PF",
    "Standard_PF",
    "catalogue_1\n PF",
    "catalogue_2\nST",
    "catalogue_3\n LIBRE",
]
cols_existantes = [c for c in cols_synthese if c in veh.index]
st.table(veh[cols_existantes].to_frame(name="Valeur"))

# ----------------- APERÇU + UPLOAD IMAGES -----------------

st.subheader("Images associées")

col1, col2, col3 = st.columns(3)

img_veh_val = veh.get("Image Vehicule")
img_client_val = veh.get("Image Client")
img_carbu_val = veh.get("Image Carburant")

img_veh_path = resolve_image_path(img_veh_val, "vehicules")
img_client_path = resolve_image_path(img_client_val, "clients")
img_carbu_path = resolve_image_path(img_carbu_val, "carburant")

with col1:
    st.write("Image véhicule (BDD)")
    uploaded_veh = st.file_uploader(
        "Remplacer (optionnel)", type=["png", "jpg", "jpeg"], key="upl_veh"
    )
    show_image(img_veh_path, "Aperçu véhicule", uploaded_file=uploaded_veh)

with col2:
    st.write("Logo client (BDD)")
    uploaded_client = st.file_uploader(
        "Remplacer (optionnel)", type=["png", "jpg", "jpeg"], key="upl_client"
    )
    show_image(img_client_path, "Aperçu client", uploaded_file=uploaded_client)

with col3:
    st.write("Picto carburant (BDD)")
    uploaded_carbu = st.file_uploader(
        "Remplacer (optionnel)", type=["png", "jpg", "jpeg"], key="upl_carbu"
    )
    show_image(img_carbu_path, "Aperçu carburant", uploaded_file=uploaded_carbu)

# ----------------- BOUTON DE TÉLÉCHARGEMENT FT -----------------

st.markdown("---")
st.subheader("Génération de la fiche technique")

ft_file = genere_ft_excel(
    veh,
    cab_prod_choice, cab_opt_choice,
    mot_prod_choice, mot_opt_choice,
    ch_prod_choice, ch_opt_choice,
    caisse_prod_choice, caisse_opt_choice,
    gf_prod_choice, gf_opt_choice,
    hay_prod_choice, hay_opt_choice,
    uploaded_veh, uploaded_client, uploaded_carbu,
)

if ft_file is not None:
    nom_fichier = f"FT_{veh.get('Code_PF', 'PF')}_{veh.get('Modele', 'MODELE')}.xlsx"
    st.download_button(
        label=" Télécharger la fiche technique remplie",
        data=ft_file,
        file_name=nom_fichier,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
else:
    st.info("Impossible de générer la fiche technique : modèle manquant ou erreur.")
